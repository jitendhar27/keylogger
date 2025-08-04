# export_logs.py
from fpdf import FPDF
from docx import Document
import sys

def export_to_pdf(log_file):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    with open(log_file, "r") as f:
        for line in f:
            pdf.cell(200, 10, txt=line, ln=True)

    pdf.output("keylog.pdf")

def export_to_word(log_file):
    doc = Document()
    
    with open(log_file, "r") as f:
        for line in f:
            doc.add_paragraph(line)

    doc.save("keylog.docx")

if __name__ == "__main__":
    log_file = sys.argv[1] if len(sys.argv) > 1 else "keylog.txt"
    export_to_pdf(log_file)
    export_to_word(log_file)
